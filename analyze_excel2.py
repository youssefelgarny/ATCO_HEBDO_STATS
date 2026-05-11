import openpyxl
import zipfile
import json
import re
from collections import defaultdict

FILE = r"C:\Users\user\Downloads\ATCO HEBDO (1) du 02-03-26 au 08-03-26 Fichier de JArrou 1.xlsx"

# Extract Power Query M code from the xlsx zip
print("=== POWER QUERY M CODE ===")
with zipfile.ZipFile(FILE, 'r') as z:
    for name in z.namelist():
        if 'query' in name.lower() or 'customXml' in name or 'connections' in name.lower():
            print(f"\n--- {name} ---")
            try:
                content = z.read(name).decode('utf-8', errors='replace')
                # Print first 3000 chars
                print(content[:3000])
            except:
                pass

# Sheet analysis
print("\n\n=== SHEET ANALYSIS ===")
wb = openpyxl.load_workbook(FILE, data_only=True)

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    state = ws.sheet_state
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    
    charts_info = []
    for chart in ws._charts:
        c = {"type": type(chart).__name__, "title": None}
        try:
            c["title"] = str(chart.title) if chart.title else None
        except: pass
        try:
            c["series"] = []
            for s in chart.series:
                try: c["series"].append(str(s.title) if s.title else str(s.val))
                except: pass
        except: pass
        charts_info.append(c)
    
    # Get first row as header
    headers = []
    try:
        first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = [h for h in first_row if h is not None][:20]
    except: pass
    
    print(f"\n[{state.upper()}] Sheet: '{sheetname}' | {max_col} cols x {max_row} rows")
    if headers:
        print(f"  Headers: {headers}")
    if charts_info:
        print(f"  Charts: {charts_info}")
